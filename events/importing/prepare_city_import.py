from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from pathlib import Path

from django.utils.text import slugify
from openpyxl import load_workbook

from events.importing.city_content import DATA_TYPES, FIELD_ALIASES, normalize_kind, normalize_text
from events.models import Category


URL_FIELDS = {
    "website",
    "source_url",
    "ticket_url",
    "reservation_url",
    "image_url",
    "venue_website",
    "organizer_website",
}

TEXT_FIELDS = {
    "city",
    "title",
    "slug",
    "kind",
    "category",
    "subcategory",
    "raw_date_text",
    "venue",
    "address",
    "source_url",
    "ticket_url",
    "venue_website",
}

VALID_KINDS = {"event", "activity", "festival", "place", "food_drink"}


@dataclass
class PreparedRow:
    row_number: int
    values: dict[str, str]


@dataclass
class PrepareIssue:
    row_number: int | None
    severity: str
    code: str
    message: str


@dataclass
class PrepareFix:
    row_number: int
    field: str
    old_value: str
    new_value: str
    reason: str


@dataclass
class PrepareResult:
    city: str
    data_type: str
    input_file: str
    output_file: str
    sheet: str
    rows: int
    blocking: list[PrepareIssue] = field(default_factory=list)
    warnings: list[PrepareIssue] = field(default_factory=list)
    info: list[PrepareIssue] = field(default_factory=list)
    fixes: list[PrepareFix] = field(default_factory=list)
    dry_run_summary: dict | None = None

    @property
    def ready_for_dry_run(self):
        return not self.blocking

    @property
    def auto_fixed_titles(self):
        return sum(1 for fix in self.fixes if fix.reason.startswith("duplicate title"))

    @property
    def normalized_url_fields(self):
        return sum(1 for fix in self.fixes if fix.field in self._url_column_names())

    def _url_column_names(self):
        return {
            alias
            for field_name in URL_FIELDS
            for alias in FIELD_ALIASES.get(field_name, [])
        }

    def to_dict(self):
        return {
            "city": self.city,
            "type": self.data_type,
            "input_file": self.input_file,
            "output_file": self.output_file,
            "sheet": self.sheet,
            "rows": self.rows,
            "ready_for_dry_run": self.ready_for_dry_run,
            "auto_fixed_titles": self.auto_fixed_titles,
            "normalized_url_fields": self.normalized_url_fields,
            "blocking": [asdict(issue) for issue in self.blocking],
            "warnings": [asdict(issue) for issue in self.warnings],
            "info": [asdict(issue) for issue in self.info],
            "fixes": [asdict(fix) for fix in self.fixes],
            "dry_run_summary": self.dry_run_summary,
        }


def prepared_output_path(path):
    source = Path(path)
    return source.with_name(f"{source.stem}_prepared{source.suffix}")


def prepare_city_import_batch(path, city_slug, data_type, output_path=None):
    if data_type not in DATA_TYPES:
        raise ValueError(f"Unsupported type '{data_type}'. Expected one of: {', '.join(sorted(DATA_TYPES))}")

    source = Path(path)
    if not source.exists():
        raise FileNotFoundError(f"Excel file not found: {source}")

    output = Path(output_path) if output_path else prepared_output_path(source)
    if source.resolve() == output.resolve():
        raise ValueError("Prepared output may not overwrite the input file.")

    workbook = load_workbook(source)
    sheet = _pick_sheet(workbook)
    headers = [normalize_cell_value(cell.value) for cell in sheet[1]]
    columns = {header: index + 1 for index, header in enumerate(headers) if header}
    field_columns = _field_columns(columns)

    result = PrepareResult(
        city=normalize_text(city_slug).casefold(),
        data_type=data_type,
        input_file=str(source),
        output_file=str(output),
        sheet=sheet.title,
        rows=max(sheet.max_row - 1, 0),
    )

    _validate_columns(result, field_columns)
    rows = _read_rows(sheet, field_columns)
    _normalize_cells(sheet, rows, field_columns, result)
    _validate_rows(rows, result, data_type)
    _auto_fix_duplicate_title_date_series(sheet, rows, field_columns, result)

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return result


def write_prepare_report(result, report_path):
    path = Path(report_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(_render_markdown_report(result), encoding="utf-8")


def write_prepare_json(result, json_path):
    path = Path(json_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(result.to_dict(), indent=2, ensure_ascii=False), encoding="utf-8")


def normalize_cell_value(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if text.casefold() in {"nan", "none", "null", "<na>", "nat"}:
        return ""
    return re.sub(r"\s+", " ", text)


def _pick_sheet(workbook):
    if "Master template" in workbook.sheetnames:
        return workbook["Master template"]
    return workbook[workbook.sheetnames[0]]


def _field_columns(columns):
    field_columns = {}
    normalized_to_actual = {column.casefold(): column for column in columns}
    for field_name, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            actual = normalized_to_actual.get(alias.casefold())
            if actual:
                field_columns[field_name] = (actual, columns[actual])
                break
    return field_columns


def _validate_columns(result, field_columns):
    required = ["title", "slug", "kind", "source_url"]
    for field_name in required:
        if field_name not in field_columns:
            result.blocking.append(
                PrepareIssue(
                    row_number=None,
                    severity="blocking",
                    code="missing_column",
                    message=f"Missing required column for {field_name}",
                )
            )


def _read_rows(sheet, field_columns):
    rows = []
    for row_number in range(2, sheet.max_row + 1):
        if all(sheet.cell(row_number, column).value is None for column in range(1, sheet.max_column + 1)):
            continue
        values = {}
        for field_name, (_, column_index) in field_columns.items():
            values[field_name] = normalize_cell_value(sheet.cell(row_number, column_index).value)
        rows.append(PreparedRow(row_number=row_number, values=values))
    return rows


def _normalize_cells(sheet, rows, field_columns, result):
    for row in rows:
        for field_name in TEXT_FIELDS | URL_FIELDS:
            if field_name not in field_columns:
                continue
            column_name, column_index = field_columns[field_name]
            old_raw = sheet.cell(row.row_number, column_index).value
            old_value = "" if old_raw is None else str(old_raw)
            new_value = normalize_cell_value(old_raw)
            if old_value != new_value:
                sheet.cell(row.row_number, column_index).value = new_value
                row.values[field_name] = new_value
                reason = "normalized URL placeholder/whitespace" if field_name in URL_FIELDS else "normalized text whitespace"
                fix = PrepareFix(row.row_number, column_name, old_value, new_value, reason)
                result.fixes.append(fix)
                result.info.append(
                    PrepareIssue(row.row_number, "info", "normalized_field", f"{column_name}: {reason}")
                )


def _validate_rows(rows, result, data_type):
    slug_counts = Counter(row.values.get("slug", "") for row in rows if row.values.get("slug", ""))
    source_url_counts = Counter(row.values.get("source_url", "") for row in rows if row.values.get("source_url", ""))
    title_venue_date_counts = Counter(_title_venue_date_key(row) for row in rows if _title_venue_date_key(row))
    title_source_counts = Counter(_title_source_key(row) for row in rows if _title_source_key(row))
    known_category_names = {normalize_text(name).casefold() for name in Category.objects.values_list("name", flat=True)}
    known_category_slugs = set(Category.objects.values_list("slug", flat=True))

    for row in rows:
        title = row.values.get("title", "")
        city = row.values.get("city", result.city)
        kind = normalize_kind(row.values.get("kind", ""), data_type)
        start_at = row.values.get("start_at", "")
        end_at = row.values.get("end_at", "")

        if not title:
            result.blocking.append(PrepareIssue(row.row_number, "blocking", "missing_title", "Missing title"))
        if not city:
            result.blocking.append(PrepareIssue(row.row_number, "blocking", "missing_city", "Missing city"))
        if not kind:
            result.blocking.append(PrepareIssue(row.row_number, "blocking", "missing_kind", "Missing kind"))
        elif kind not in VALID_KINDS:
            result.blocking.append(PrepareIssue(row.row_number, "blocking", "invalid_kind", f"Invalid kind: {kind}"))

        slug = row.values.get("slug", "")
        if slug and slug_counts[slug] > 1:
            result.blocking.append(PrepareIssue(row.row_number, "blocking", "duplicate_slug", f"Duplicate slug: {slug}"))

        tvd_key = _title_venue_date_key(row)
        if tvd_key and title_venue_date_counts[tvd_key] > 1:
            result.blocking.append(
                PrepareIssue(row.row_number, "blocking", "duplicate_title_venue_date", "Same title + venue + date")
            )

        source_key = _title_source_key(row)
        if source_key and title_source_counts[source_key] > 1 and not _has_distinct_date_series(rows, source_key):
            result.blocking.append(
                PrepareIssue(row.row_number, "blocking", "duplicate_title_source_url", "Same title + source_url")
            )

        if _end_before_start(start_at, end_at):
            result.blocking.append(PrepareIssue(row.row_number, "blocking", "end_before_start", "EndAt is before StartAt"))

        _append_warnings(row, result, source_url_counts, known_category_names, known_category_slugs)


def _append_warnings(row, result, source_url_counts, known_category_names, known_category_slugs):
    category = normalize_text(row.values.get("category", ""))
    source_url = normalize_text(row.values.get("source_url", ""))
    if category and category.casefold() not in known_category_names and slugify(category) not in known_category_slugs:
        result.warnings.append(PrepareIssue(row.row_number, "warning", "unknown_category", f"Unknown category: {category}"))
    if source_url and source_url_counts[source_url] > 1:
        result.warnings.append(PrepareIssue(row.row_number, "warning", "shared_source_url", "Shared source_url"))
    if not row.values.get("image_url", ""):
        result.warnings.append(PrepareIssue(row.row_number, "warning", "missing_image_url", "Missing image_url"))
    if not row.values.get("ticket_url", ""):
        result.warnings.append(PrepareIssue(row.row_number, "warning", "missing_ticket_url", "Missing ticket_url"))
    if not row.values.get("price_note", ""):
        result.warnings.append(PrepareIssue(row.row_number, "warning", "missing_price_note", "Missing price_note"))
    if not row.values.get("latitude", "") or not row.values.get("longitude", ""):
        result.warnings.append(PrepareIssue(row.row_number, "warning", "missing_coordinates", "Missing coordinates"))


def _auto_fix_duplicate_title_date_series(sheet, rows, field_columns, result):
    if "title" not in field_columns:
        return
    title_column_name, title_column_index = field_columns["title"]
    groups = defaultdict(list)
    for row in rows:
        title = row.values.get("title", "")
        if title:
            groups[_key(title)].append(row)

    for grouped_rows in groups.values():
        if len(grouped_rows) <= 1:
            continue
        if not _is_auto_fixable_date_series(grouped_rows):
            continue
        for row in grouped_rows:
            old_title = row.values["title"]
            date_text = row.values.get("raw_date_text") or row.values.get("start_at")
            new_title = f"{old_title} - {date_text}"
            sheet.cell(row.row_number, title_column_index).value = new_title
            row.values["title"] = new_title
            result.fixes.append(
                PrepareFix(
                    row_number=row.row_number,
                    field=title_column_name,
                    old_value=old_title,
                    new_value=new_title,
                    reason="duplicate title date series auto-fixed with existing date",
                )
            )
            result.info.append(
                PrepareIssue(row.row_number, "info", "auto_fixed_title", f"Duplicate title made distinct with date: {date_text}")
            )


def _is_auto_fixable_date_series(rows):
    slugs = [row.values.get("slug", "") for row in rows]
    dates = [row.values.get("raw_date_text") or row.values.get("start_at") for row in rows]
    if len(set(slugs)) != len(slugs):
        return False
    if not all(dates):
        return False
    if len(set(dates)) != len(dates):
        return False
    title_venue_dates = [_title_venue_date_key(row) for row in rows]
    return len(set(title_venue_dates)) == len(title_venue_dates)


def _key(value):
    return " ".join(normalize_text(value).casefold().split())


def _title_venue_date_key(row):
    title = _key(row.values.get("title", ""))
    venue = _key(row.values.get("venue", ""))
    date_value = _key(row.values.get("raw_date_text", "") or row.values.get("start_at", ""))
    if not title or not venue or not date_value:
        return ""
    return title, venue, date_value


def _title_source_key(row):
    title = _key(row.values.get("title", ""))
    source_url = _key(row.values.get("source_url", ""))
    if not title or not source_url:
        return ""
    return title, source_url


def _has_distinct_date_series(rows, source_key):
    matching = [row for row in rows if _title_source_key(row) == source_key]
    dates = [row.values.get("raw_date_text") or row.values.get("start_at") for row in matching]
    slugs = [row.values.get("slug", "") for row in matching]
    return len(matching) > 1 and all(dates) and len(set(dates)) == len(dates) and len(set(slugs)) == len(slugs)


def _end_before_start(start_at, end_at):
    start = _parse_datetime(start_at)
    end = _parse_datetime(end_at)
    return bool(start and end and end < start)


def _parse_datetime(value):
    text = normalize_text(value)
    if not text:
        return None
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        return None


def _render_markdown_report(result):
    lines = [
        "# City Import Batch Preparation",
        "",
        f"City: {result.city}",
        f"Type: {result.data_type}",
        f"Rows: {result.rows}",
        f"Input: `{result.input_file}`",
        f"Cleaned file: `{result.output_file}`",
        "",
        "## Summary",
        "",
        f"- Blocking issues: {len(result.blocking)}",
        f"- Warnings: {len(result.warnings)}",
        f"- Auto-fixed titles: {result.auto_fixed_titles}",
        f"- Normalized URL fields: {result.normalized_url_fields}",
        f"- Ready for dry-run: {'yes' if result.ready_for_dry_run else 'no'}",
    ]
    if result.blocking:
        lines.extend(["", "## Blocking"])
        for issue in result.blocking[:50]:
            lines.append(f"- row {issue.row_number}: {issue.code} - {issue.message}")
    if result.fixes:
        lines.extend(["", "## Fixes"])
        for fix in result.fixes[:50]:
            lines.append(f"- row {fix.row_number}: {fix.field} -> `{fix.new_value}`")
    return "\n".join(lines) + "\n"
